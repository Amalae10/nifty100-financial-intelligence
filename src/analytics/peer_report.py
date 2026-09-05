import os
import sqlite3
import pandas as pd
from openpyxl.styles import PatternFill, Font

DB = "nifty100.db"
OUTPUT = "output/peer_comparison.xlsx"

GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="FFC7CE")
GOLD = PatternFill("solid", fgColor="FFD966")

METRICS = {
    "ROE": "return_on_equity_pct",
    "ROCE": "return_on_capital_pct",
    "NPM": "net_profit_margin_pct",
    "OPM": "operating_profit_margin_pct",
    "D/E": "debt_to_equity",
    "Interest Coverage": "interest_coverage",
    "FCF": "free_cash_flow_cr",
    "FCF CAGR 5Y": "fcf_cagr_5yr",
    "CFO/PAT": "cfo_pat_ratio",
    "FCF Conversion": "fcf_conversion_pct",
    "Asset Turnover": "asset_turnover",
    "Revenue CAGR 3Y": "revenue_cagr_3yr",
    "Revenue CAGR 5Y": "revenue_cagr_5yr",
    "PAT CAGR 3Y": "pat_cagr_3yr",
    "PAT CAGR 5Y": "pat_cagr_5yr",
    "EPS CAGR 3Y": "eps_cagr_3yr",
    "EPS CAGR 5Y": "eps_cagr_5yr",
    "P/E": "pe_ratio",
    "Dividend Yield": "dividend_yield_pct",
    "Composite Score": "composite_quality_score",
}


def load_data():
    with sqlite3.connect(DB) as conn:
        ratios = pd.read_sql(
            "SELECT * FROM financial_ratios", conn
        )

        peers = pd.read_sql(
            "SELECT * FROM peer_groups", conn
        )

        companies = pd.read_sql(
            "SELECT id, company_name FROM companies", conn
        )

    companies = companies.rename(
        columns={"id": "company_id"}
    )

    return ratios, peers, companies


def latest_data(df):
    df = df[df["year"] != "TTM"].copy()

    df["year_num"] = pd.to_numeric(
        df["year"]
        .astype(str)
        .str.extract(r"(\d{4})")[0],
        errors="coerce"
    )

    return (
        df.sort_values("year_num")
        .groupby("company_id")
        .tail(1)
        .copy()
    )


def percent_rank(series, inverse=False):
    valid = series.dropna()
    result = pd.Series(float("nan"), index=series.index)

    if len(valid) <= 1:
        result.loc[valid.index] = 0.0
        return result

    ranks = valid.rank(method="min")
    pct = (ranks - 1) / (len(valid) - 1)

    if inverse:
        pct = 1 - pct

    result.loc[valid.index] = pct
    return result


def generate_report():
    os.makedirs("output", exist_ok=True)

    ratios, peers, companies = load_data()
    ratios = latest_data(ratios)

    data = peers.merge(
        companies,
        on="company_id",
        how="left"
    )

    data = data.merge(
        ratios,
        on="company_id",
        how="left"
    )

    with pd.ExcelWriter(
        OUTPUT,
        engine="openpyxl"
    ) as writer:

        for group_name, group in data.groupby(
            "peer_group_name"
        ):

            group = group.copy()

            out = pd.DataFrame({
                "company_id": group["company_id"],
                "company_name": group["company_name"]
            })

            pct_cols = []

            for label, col in METRICS.items():

                out[label] = group[col]

                pct_col = f"{label} Percentile"

                out[pct_col] = percent_rank(
                    group[col],
                    inverse=(label == "D/E")
                )

                pct_cols.append(pct_col)

            median = {
                "company_id": "MEDIAN",
                "company_name": "Peer Group Median"
            }

            for label in METRICS:
                median[label] = out[label].median()
                median[f"{label} Percentile"] = None

            out = pd.concat(
                [out, pd.DataFrame([median])],
                ignore_index=True
            )

            sheet_name = group_name[:31]

            out.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

            ws = writer.book[sheet_name]

            headers = {
                cell.value: cell.column
                for cell in ws[1]
            }

            # Percentile colours
            for pct_col in pct_cols:

                col_num = headers[pct_col]

                for row in range(2, ws.max_row):

                    cell = ws.cell(
                        row=row,
                        column=col_num
                    )

                    value = cell.value

                    if not isinstance(
                        value,
                        (int, float)
                    ):
                        continue

                    if value >= 0.75:
                        cell.fill = GREEN

                    elif value <= 0.25:
                        cell.fill = RED

                    else:
                        cell.fill = YELLOW

            # Benchmark company row
            benchmark_ids = group[
                group["is_benchmark"] == 1
            ]["company_id"].tolist()

            for row in range(2, ws.max_row):

                company_id = ws.cell(
                    row=row,
                    column=1
                ).value

                if company_id in benchmark_ids:

                    for cell in ws[row]:
                        cell.fill = GOLD
                        cell.font = Font(
                            bold=True
                        )

            # Median row formatting
            for cell in ws[ws.max_row]:
                cell.font = Font(
                    bold=True
                )

            # Freeze company columns
            ws.freeze_panes = "C2"

            # Auto column width
            for column in ws.columns:

                max_len = max(
                    len(str(cell.value))
                    if cell.value is not None
                    else 0
                    for cell in column
                )

                col_letter = (
                    column[0].column_letter
                )

                ws.column_dimensions[
                    col_letter
                ].width = min(
                    max_len + 2,
                    22
                )

    print("created:", OUTPUT)


if __name__ == "__main__":
    generate_report()