import os
import pandas as pd
from openpyxl.styles import PatternFill

from src.screener.engine import (
    run_preset, load_data, latest_data, load_config)

from src.analytics.composite_score import (add_composite_score,sector_relative_score)

OUTPUT = "output/screener_output.xlsx"

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")

PRESETS = [
    "quality_compounder",
    "value_pick",
    "growth_accelerator",
    "dividend_champion",
    "debt_free_blue_chip",
    "turnaround_watch"
]

KPI_COLUMNS = [
    "company_id", "year", "broad_sector",
    "composite_quality_score", "sector_composite_score",
    "return_on_equity_pct", "return_on_capital_pct",
    "net_profit_margin_pct", "free_cash_flow_cr",
    "fcf_cagr_5yr", "cfo_pat_ratio",
    "revenue_cagr_3yr", "revenue_cagr_5yr",
    "pat_cagr_5yr", "eps_cagr_5yr",
    "debt_to_equity", "interest_coverage",
    "pe_ratio", "pb_ratio", "dividend_yield_pct"
]


def colour_threshold(writer, sheet, filters):
    ws = writer.book[sheet]
    headers = {c.value: c.column for c in ws[1]}
    metrics = load_config()["metrics"]

    for rule, limit in filters.items():
        col = metrics.get(rule)

        if col not in headers:
            continue

        for cell in list(ws.columns)[headers[col] - 1][1:]:

            if cell.value is None:
                continue

            if rule == "de_eq":
                passed = cell.value == limit

            elif rule.endswith("_min"):
                passed = cell.value > limit

            elif rule.endswith("_max"):
                passed = cell.value < limit

            else:
                continue

            cell.fill = GREEN if passed else RED

def export_screeners():
    os.makedirs("output", exist_ok=True)

    # Sector score using full universe
    full_df = latest_data(load_data())
    full_df = add_composite_score(full_df)
    full_df = sector_relative_score(full_df)

    score_map = full_df[
    [
        "company_id",
        "composite_quality_score",
        "sector_composite_score"
    ]
]

    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:

        for preset in PRESETS:
            df = run_preset(preset)

            df = df.drop(
                columns=[
                "composite_quality_score",
                "sector_composite_score"],
                errors="ignore")

            df = df.merge(
                score_map,
                on="company_id",
                how="left")

            df = df.sort_values(
                "composite_quality_score",
                ascending=False
            )

            df = df[
                [c for c in KPI_COLUMNS if c in df.columns]
            ]

            sheet = preset[:31]

            df.to_excel(
                writer,
                sheet_name=sheet,
                index=False
            )

            colour_threshold(
                writer,
                sheet,
                load_config()["presets"][preset]
            )

    print("created:", OUTPUT)


if __name__ == "__main__":
    export_screeners()